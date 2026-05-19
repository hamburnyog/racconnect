import csv
import openpyxl
from datetime import datetime

def get_id_mapping():
    # Load mapping from R4A sheet
    file_path = 'BIO ID OF RACC EMPLOYEES.xlsx'
    wb = openpyxl.load_workbook(file_path, data_only=True)
    if 'R4A' not in wb.sheetnames:
        print("Error: Sheet 'R4A' not found.")
        return {}
    
    sheet = wb['R4A']
    mapping = {}
    data = list(sheet.iter_rows(values_only=True))
    
    # Identify header indices
    co_idx = -1
    dev_idx = -1
    
    for i, row in enumerate(data):
        row_values = [str(c).strip() if c is not None else "" for c in row]
        if 'BIO ID NO. (from CO)' in row_values:
            co_idx = row_values.index('BIO ID NO. (from CO)')
            dev_idx = row_values.index('BIO ID NO. (from device)')
            data_start = i + 1
            break
    else:
        # Defaults based on manual inspection if headers not found
        co_idx = 5
        dev_idx = 6
        data_start = 2

    for row in data[data_start:]:
        if len(row) <= max(co_idx, dev_idx):
            continue
        co_id = row[co_idx]
        device_id = row[dev_idx]
        if co_id is not None and device_id is not None:
            try:
                # Map raw device ID (int) to CO ID (str)
                mapping[int(device_id)] = str(co_id)
            except (ValueError, TypeError):
                pass
    
    # Manual overrides/additions provided by user
    mapping[783] = "14035"
    mapping[820] = "14040"
    
    return mapping

def process_logs():
    mapping = get_id_mapping()
    
    output_rows = []
    
    # Filter range: May 1-15, 2026
    start_date = datetime(2026, 5, 1)
    end_date = datetime(2026, 5, 15, 23, 59, 59)
    
    try:
        with open('record.txt', 'r') as f:
            for line in f:
                parts = line.strip().split(',')
                if len(parts) != 3:
                    continue
                
                raw_id, date_str, time_str = parts
                
                try:
                    date_obj = datetime.strptime(date_str, '%m/%d/%Y')
                    if not (start_date <= date_obj <= end_date):
                        continue
                    
                    # Get mapped ID
                    device_id = int(raw_id)
                    if device_id not in mapping:
                        # Optional: print or log missing mappings
                        continue
                        
                    co_id = mapping[device_id]
                    # Format: MM/DD/YYYY HH:MM:SS
                    timestamp = f"{date_str} {time_str}"
                    output_rows.append([co_id, timestamp])
                except (ValueError, TypeError):
                    continue
    except FileNotFoundError:
        print("Error: record.txt not found")
        return

    # Sort output by ID then Timestamp
    output_rows.sort(key=lambda x: (x[0], datetime.strptime(x[1], '%m/%d/%Y %H:%M:%S')))

    output_file = 'upload_time_logs_may_2026.csv'
    with open(output_file, 'w', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['bio_id', 'time_logs'])
        writer.writerows(output_rows)
    
    print(f"Processed {len(output_rows)} log entries into {output_file}")

if __name__ == '__main__':
    process_logs()
